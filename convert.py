import sys,re

def pdf_to_word(src,dst):
    from pdf2docx import Converter
    cv=Converter(src)
    try: cv.convert(dst)
    finally: cv.close()

def clean_font(n):
    n=(n or 'Arial').split('+')[-1]
    n=re.sub(r'[-,_](Bold|Italic|Roman|Regular|MT|PS)$','',n,flags=re.I)
    return n[:31] or 'Arial'

def cluster(vals,tol):
    groups=[]
    for v in sorted(vals):
        if not groups or abs(v-groups[-1][-1])>tol: groups.append([v])
        else: groups[-1].append(v)
    return [sum(g)/len(g) for g in groups]

def pdf_to_excel(src,dst):
    import fitz,pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font,Border,Side,Alignment,PatternFill
    from openpyxl.utils import get_column_letter
    wb=Workbook(); wb.remove(wb.active); thin=Side(style='thin',color='B7BEC9'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
    doc=fitz.open(src)
    try:
      with pdfplumber.open(src) as pl:
        for pi,page in enumerate(doc):
          ws=wb.create_sheet(f'Page {pi+1}'); tables=[]
          try: tables=pl.pages[pi].extract_tables(table_settings={'vertical_strategy':'lines','horizontal_strategy':'lines','intersection_tolerance':5,'snap_tolerance':4}) or []
          except Exception: pass
          if tables:
            row0=1
            for table in tables:
              for row in table:
                for ci,val in enumerate(row or [],1):
                  c=ws.cell(row0,ci,val or ''); c.border=border; c.alignment=Alignment(vertical='top',wrap_text=True)
                row0+=1
              row0+=1
          else:
            words=page.get_text('words'); ys=cluster([w[1] for w in words],6) if words else []; xs=cluster([w[0] for w in words],14) if words else []
            for w in words:
              x0,y0,x1,y1,text=w[:5]; r=min(range(len(ys)),key=lambda i:abs(ys[i]-y0))+1; c=min(range(len(xs)),key=lambda i:abs(xs[i]-x0))+1
              cell=ws.cell(r,c); cell.value=(str(cell.value)+' '+text).strip() if cell.value else text; cell.border=border; cell.alignment=Alignment(vertical='top',wrap_text=True)
            try:
              blocks=page.get_text('dict').get('blocks',[])
              for block in blocks:
                for line in block.get('lines',[]):
                  for sp in line.get('spans',[]):
                    if not sp.get('text','').strip() or not ys or not xs: continue
                    r=min(range(len(ys)),key=lambda i:abs(ys[i]-sp['bbox'][1]))+1; c=min(range(len(xs)),key=lambda i:abs(xs[i]-sp['bbox'][0]))+1
                    cell=ws.cell(r,c); flags=sp.get('flags',0); cell.font=Font(name=clean_font(sp.get('font')),size=max(8,min(28,float(sp.get('size',10)))),bold=bool(flags&16),italic=bool(flags&2)); break
            except Exception: pass
          for ci in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(ci)].width=18
          ws.sheet_view.showGridLines=False; ws.freeze_panes='A1'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    finally: doc.close()
    wb.save(dst)

if __name__=='__main__':
    kind,src,dst=sys.argv[1:]
    if kind=='word': pdf_to_word(src,dst)
    elif kind=='excel': pdf_to_excel(src,dst)
    else: raise SystemExit('Unknown conversion')
